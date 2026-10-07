#!/usr/bin/env python3
"""Build reviewable manuscript DOCX and publication artwork from frozen evidence."""
from pathlib import Path
import json,csv,re,shutil
from datetime import datetime
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
REPO=Path(__file__).resolve().parents[1];OUT=REPO/'paper/submission';FIG=OUT/'figures';SUP=OUT/'supplement'
plt.rcParams.update({'font.family':'DejaVu Serif','font.size':10,'pdf.fonttype':42,'ps.fonttype':42})
figure_words=[]
def save(fig,name):
 fig.canvas.draw()
 figure_words.extend(t.get_text() for t in fig.findobj(matplotlib.text.Text) if t.get_visible())
 fig.savefig(FIG/(name+'.png'),dpi=400,bbox_inches='tight',facecolor='white')
 fig.savefig(FIG/(name+'.pdf'),bbox_inches='tight',facecolor='white')
 plt.close(fig)
def figures():
 FIG.mkdir(exist_ok=True)
 s=json.loads((SUP/'results_summary.json').read_text());pair_n=s['different_card_pairs'];cost_n=s['conversion_included']
 fig,ax=plt.subplots(figsize=(8,4.3));ax.set_xlim(0,10);ax.set_ylim(0,5);ax.axis('off')
 def box(x,y,w,h,text):
  ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=0.08',facecolor='#f4f4f4',edgecolor='black'))
  ax.text(x+w/2,y+h/2,text,ha='center',va='center',fontsize=10)
 box(.2,2,2,1,'Collection\ncontains A\nCapabilities at t₀')
 box(4.2,3.4,5.2,1.1,'Revision of A\nNew capabilities attached to existing card')
 box(4.2,.6,5.2,1.8,'Acquire different card B\nKeep A: C(B)\nEligible surrender: net C(B) − D(A)\nAdditional: max(0, net)\nSurplus: max(0, −net)')
 for xy in [(4.1,3.95),(4.1,1.5)]:ax.annotate('',xy=xy,xytext=(2.3,2.5),arrowprops={'arrowstyle':'->','lw':1.3})
 ax.text(2.8,3.5,'Update',ha='center');ax.text(2.7,1.45,'Acquisition',ha='center')
 ax.text(5,.1,'Nominal tier comparison: C(B) − C(A)     •     Core/free A: D(A) = 0, A remains held',ha='center',fontsize=9)
 save(fig,'figure_1')
 fig,axes=plt.subplots(1,2,figsize=(9,4.5),gridspec_kw={'width_ratios':[1,1.6]})
 vals=[s['rarity_transitions'][k] for k in ['increased','unchanged','decreased']]
 bars=axes[0].bar(['Higher','Equal','Lower'],vals,color=['#555555','#aaaaaa','#eeeeee'],edgecolor='black',hatch='')
 for b,n in zip(bars,vals):axes[0].text(b.get_x()+b.get_width()/2,n+1,f'{n}/{pair_n}\n({n/pair_n:.1%})',ha='center',fontsize=9)
 axes[0].set_ylim(0,85);axes[0].set_ylabel('Directed different-card identity pairs');axes[0].set_title(f'A. Candidate rarity (n = {pair_n})',loc='left',fontsize=11)
 for key,label,style in [('nominal_price_difference','Nominal difference','-'),('candidate_craft','Acquire B, keep A','--'),('additional_dust','Additional conversion dust',':')]:
  ds=s['matched_sample'][key]['distribution'];counts=sorted((int(x),n) for x,n in ds.items());xx=[-320];yy=[0];c=0
  for x,n in counts:c+=n;xx.append(x);yy.append(c/cost_n)
  xx.append(420);yy.append(1);axes[1].step(xx,yy,where='post',label=label,linestyle=style,color='black',linewidth=1.6)
 axes[1].axvline(0,color='.7',lw=.7);axes[1].set_xlim(-320,420);axes[1].set_ylim(0,1.04);axes[1].set_xlabel('Dust per single-copy scenario');axes[1].set_ylabel(f'Cumulative proportion of {cost_n} eligible pairs');axes[1].set_title(f'B. Matched resource sample (n = {cost_n})',loc='left',fontsize=11);axes[1].legend(loc='lower right',fontsize=8,frameon=False)
 fig.text(.5,.015,f"Coverage: {cost_n}/{pair_n} included; {s['cost_exclusions']['Candidate has no ordinary crafting route']} no ordinary crafting route; {s['cost_exclusions']['Unresolved acquisition status']} acquisition-uncertain. Incomplete-effect comparisons excluded.",ha='center',fontsize=8)
 fig.tight_layout(rect=[0,.04,1,1]);save(fig,'figure_2')
 events=list(csv.DictReader((SUP/'case_events.csv').open()));ys={'Tyrantus':3,'The Demon Seed':2,'Big Game Hunter':1,'Evasive Wyrm':0}
 fig,ax=plt.subplots(figsize=(8,6.7));markers={'buff':'^','restriction':'v','repair':'s','wording':'o'}
 labels={('Tyrantus','21.8'):'Repair\n21.8',('Tyrantus','27.6'):'14/14 + Taunt\n27.6',('Tyrantus','29.0'):'Elusive wording\n29.0',('The Demon Seed','21.3'):'8/8/8\nWild ban\n21.3',('The Demon Seed','26.0'):'10/10/10\nWild unban\n26.0',('The Demon Seed','31.2.2'):'12/12/12\n31.2.2',('Big Game Hunter','5.0'):'3 → 5 Mana\n5.0 notes*',('Big Game Hunter','20.0'):'5 → 4 Mana\n20.0',('Big Game Hunter','29.0'):'Tradeable\n29.0',('Evasive Wyrm','29.0'):'Equivalent Elusive\n29.0',('Evasive Wyrm','32.0'):'5/3 → 5/4\n32.0'}
 for card,y in ys.items():ax.axhline(y,color='.8',lw=.7)
 for e in events:
  x=datetime.fromisoformat(e['event_date']);y=ys[e['card']];hollow='Source-history' in e['date_basis']
  ax.scatter(x,y,marker=markers[e['event_kind']],s=75,facecolors='white' if hollow else '#555555',edgecolors='black',zorder=3)
  key=(e['card'],e['patch']);below=key in [('Tyrantus','29.0'),('Evasive Wyrm','32.0')]
  ax.annotate(labels[key],(x,y),xytext=(0,-20 if below else 15),textcoords='offset points',ha='center',va='top' if below else 'bottom',fontsize=10,arrowprops={'arrowstyle':'-','color':'.5','lw':.6})
  if 'Two-week' in e['refund']:
   ax.text(x,y-.13,'R: 2 weeks',fontsize=7,ha='center',va='top')
  elif 'Limited full' in e['refund']:ax.text(x,y-.13,'R: limited',fontsize=7,ha='center',va='top')
 ax.set_yticks(list(ys.values()),list(ys));ax.set_ylim(-.7,3.6);ax.set_xlim(datetime(2015,7,1),datetime(2026,1,1));ax.xaxis.set_major_locator(mdates.YearLocator());ax.xaxis.set_major_formatter(mdates.DateFormatter('%Y'));ax.set_xlabel('Verified launch date or source-linked patch schedule (* = notes publication date)')
 for kind,m in markers.items():ax.scatter([],[],marker=m,color='#555555',label=kind.capitalize())
 ax.legend(loc='upper left',ncol=4,fontsize=8,frameon=False)
 ax.spines[['top','right','left']].set_visible(False);ax.tick_params(axis='y',length=0)
 fig.text(.5,.035,'Hollow marks: source-history transition + launch schedule. R: stated refund period; exact cutoffs not reconstructed.',ha='center',fontsize=8)
 fig.text(.5,.01,'Core access: Big Game Hunter from 30 Mar 2021; 2024 rotation 19 Mar; 2025 rotation 25 Mar. Core copies return no dust.',ha='center',fontsize=8)
 fig.tight_layout(rect=[0,.065,1,1]);save(fig,'figure_3')
def format_doc(doc):
 sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
 sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
 for st in doc.styles:
  if st.type==1:
   st.font.name='Times New Roman';st.font.size=Pt(12)
   st.paragraph_format.line_spacing=2;st.paragraph_format.space_after=Pt(0)
 doc.core_properties.author='';doc.core_properties.last_modified_by='';doc.core_properties.title='Maintaining the Game, Maintaining the Collection'
 footer=sec.footer.paragraphs[0];footer.alignment=2
 fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
def inline(p,text):
 for i,part in enumerate(re.split(r'(\*[^*]+\*)',text)):
  r=p.add_run(part.strip('*') if part.startswith('*') else part)
  if part.startswith('*'):r.italic=True
  r.font.name='Times New Roman';r.font.size=Pt(12)
def convert(src,dst):
 doc=Document();format_doc(doc);lines=src.read_text().splitlines();i=0
 while i<len(lines):
  line=lines[i]
  if not line.strip():i+=1;continue
  if line.startswith('|'):
   rows=[]
   while i<len(lines) and lines[i].startswith('|'):
    cells=[x.strip() for x in lines[i].strip('|').split('|')]
    if not all(re.fullmatch(r':?-+:?',c) for c in cells):rows.append(cells)
    i+=1
   table=doc.add_table(rows=0,cols=len(rows[0]));table.style='Table Grid'
   for j,row in enumerate(rows):
    cells=table.add_row().cells
    for c,text in zip(cells,row):inline(c.paragraphs[0],text)
    if j==0:
     for c in cells:
      for r in c.paragraphs[0].runs:r.bold=True
   continue
  if line.startswith('!['):
   m=re.match(r'!\[(.*?)\]\((.*?)\)',line)
   doc.add_picture(str(src.parent/m.group(2)),width=Inches(6.5))
   pic=doc.inline_shapes[-1]._inline.docPr;pic.set('descr',m.group(1));i+=1;continue
  if line.startswith('#'):
   level=len(line)-len(line.lstrip('#'));title=line[level:].strip()
   if title=='Tables and figures':doc.add_page_break()
   doc.add_heading(title,level=min(level,3));i+=1;continue
  p=doc.add_paragraph();inline(p,line);i+=1
 doc.save(dst)
 return doc
def main():
 figures()
 doc=convert(OUT/'manuscript.md',OUT/'manuscript.docx')
 convert(OUT/'title_page.md',OUT/'title_page.docx')
 text=' '.join(p.text for p in doc.paragraphs)+' '+ ' '.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
 abstract=(OUT/'manuscript.md').read_text().split('## Abstract')[1].split('Keywords:')[0]
 count=lambda t:len(re.findall(r"\b[\w]+(?:[’'−–-][\w]+)*\b",t))
 audit={'manuscript_text_including_tables_and_captions':count(text),'figure_text_conservative_extra':count(' '.join(figure_words)),'all_inclusive_conservative_word_count':count(text)+count(' '.join(figure_words)),'abstract_words':count(abstract),'keywords':5,'format':'US Letter, 1-inch margins, 12pt Times New Roman, double spacing; artwork/tables at end','anonymous_metadata_author':doc.core_properties.author,'author_declarations':'Pending author-provided facts; see title_page.md'}
 assert audit['all_inclusive_conservative_word_count']<=7800
 assert audit['abstract_words']<=150
 assert not any(x in text for x in ['Jan Sawicki','Jack Black','pw.edu.pl','/mnt/','/home2/'])
 (OUT/'editorial_checks.json').write_text(json.dumps(audit,indent=2)+'\n')
 print(json.dumps(audit,indent=2))
if __name__=='__main__':main()
